"""
Pega la asistencia y devocionales que cada tutora registra manualmente
en su Drive sobre el reporte "unido" (notas + checks de Moodle),
calcula la nota de asistencia, compara Drive vs Aula Virtual, y da
formato final al Excel — incluyendo el resaltado de APROBADO (amarillo)
/ DESAPROBADO (rojo, letras blancas) que pide el punto 5 del pedido original.

DISEÑO: todas las columnas del reporte principal se ubican buscando su
ENCABEZADO (fila 1) por nombre, nunca por número de columna fijo. Así,
si un curso no tiene TF o EF, esas columnas simplemente no existen y el
resto del reporte no se desalinea (antes, una posición fija como
`col_ef_drive = 11` se corría entera apenas cambiaba una columna
anterior — esto es lo que el punto 2 y 3 del pedido describían).

⚠️ Los offsets de lectura del archivo de la TUTORA (fila de inicio,
columna del DNI, rango de asistencia, columna de EF, rango de
devocionales) se preservaron tal cual estaban en el código original
porque no hay un archivo de ejemplo de una tutora para verificarlos.
Viven en config.py (TUTOR_*), listos para ajustar si no calzan.
"""
from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
from openpyxl.worksheet.worksheet import Worksheet

import config
from domain.curso import ConfiguracionCurso
from services.utils import autoajustar_columnas, set_fill_por_valor, contar_celdas_con_valor, normalizar_dni_valor

logger = logging.getLogger(__name__)


class ReporteFinalError(Exception):
    pass


@dataclass
class ResultadoCalculo:
    dnis_no_encontrados_en_reporte: list[str] = field(default_factory=list)
    tutores_procesados: list[str] = field(default_factory=list)
    archivos_tutor_sin_hoja_asistencia: list[str] = field(default_factory=list)


BORDE_FINO = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


# ============================================================
# Utilidades para ubicar columnas por nombre de encabezado
# ============================================================

def _mapa_encabezados(ws: Worksheet) -> dict[str, int]:
    mapa = {}
    for col in range(1, ws.max_column + 1):
        valor = ws.cell(1, col).value
        if valor:
            mapa[str(valor).strip()] = col
    return mapa


def _col(mapa: dict[str, int], nombre: str, obligatoria: bool = True) -> int | None:
    if nombre in mapa:
        return mapa[nombre]
    if obligatoria:
        raise ReporteFinalError(
            f"No se encontró la columna '{nombre}' en el reporte unido. "
            f"Columnas disponibles: {list(mapa.keys())}. Esto normalmente indica que "
            "reporte_unido no generó esa columna: revisa los pasos previos."
        )
    return None


def _agregar_columna(ws: Worksheet, titulo: str, mapa: dict[str, int], columnas_estilizadas: set[int] | None = None) -> int:
    """Agrega una columna nueva al final de la hoja y la registra en el mapa.
    Si se pasa `columnas_estilizadas`, se anota el índice ahí para que
    _aplicar_formato_general() sepa que este encabezado ya tiene su
    propio estilo y no debe tocarlo (ver esa función para el porqué)."""
    col = ws.max_column + 1
    cell = ws.cell(1, col, titulo)
    cell.font = Font(bold=True, color=config.COLOR_BLANCO)
    cell.fill = PatternFill("solid", fgColor=config.COLOR_MORADO_HEADER)
    mapa[titulo] = col
    if columnas_estilizadas is not None:
        columnas_estilizadas.add(col)
    return col


def _a_entero(valor) -> int:
    try:
        return round(float(str(valor).strip()))
    except (TypeError, ValueError):
        return 0


# ============================================================
# Paso 1: preparar columnas nuevas en el reporte principal
# ============================================================

def _preparar_columnas(ws: Worksheet, curso: ConfiguracionCurso) -> dict:
    """Agrega todas las columnas calculadas al final del reporte y
    devuelve un dict con los índices de columna relevantes."""
    mapa = _mapa_encabezados(ws)
    columnas_estilizadas: set[int] = set()

    idx = {
        "dni": _col(mapa, "DNI"),
        "nombre": _col(mapa, "Nombre"),
        "asistencia_av": _col(mapa, "Asistencia AV"),
        "ef_av": _col(mapa, "EF (AV)", obligatoria=False),
        "tf_av": _col(mapa, "TF (AV)", obligatoria=False),
    }

    # "Grupo" no se usa para mostrar el grupo de Moodle: se llena con el
    # nombre del archivo de la tutora que trajo cada alumno, para poder
    # rastrear de qué Drive salió cada fila. Si ya existía (por venir del
    # export de notas), se reutiliza la misma columna; si no, se crea.
    idx["grupo"] = _col(mapa, "Grupo", obligatoria=False)
    if not idx["grupo"]:
        idx["grupo"] = _agregar_columna(ws, "Grupo", mapa, columnas_estilizadas)

    idx["clases_drive"] = _agregar_columna(ws, "Clases Drive", mapa, columnas_estilizadas)
    idx["asistencia_drive"] = _agregar_columna(ws, "Asistencia Drive", mapa, columnas_estilizadas)

    if curso.tiene_examen_final:
        idx["ef_drive"] = _agregar_columna(ws, "EF Drive", mapa, columnas_estilizadas)
    else:
        idx["ef_drive"] = None

    if curso.tiene_trabajo_final:
        idx["tf_drive"] = _agregar_columna(ws, "TF Drive", mapa, columnas_estilizadas)
    else:
        idx["tf_drive"] = None

    # Devocionales / % Devocionales se calculan siempre, para todo curso.
    idx["devo_total"] = _agregar_columna(ws, "Devocionales Entregados", mapa, columnas_estilizadas)
    idx["devo_pct"] = _agregar_columna(ws, "% Devocionales", mapa, columnas_estilizadas)

    idx["estado_final"] = _agregar_columna(ws, "Estado Final", mapa, columnas_estilizadas)

    # Columnas S1..S{num_clases}: asistencia por clase copiada del Drive
    idx["clases_drive_inicio"] = ws.max_column + 1
    for i in range(1, curso.num_clases + 1):
        _agregar_columna(ws, f"S{i}", mapa, columnas_estilizadas)

    idx["mapa"] = mapa
    idx["columnas_estilizadas"] = columnas_estilizadas
    return idx


# ============================================================
# Paso 2: nota de asistencia (0-20) según clases asistidas
# ============================================================

def calcular_nota_asistencia(clases_asistidas: int, curso: ConfiguracionCurso) -> int:
    if curso.num_clases == 9:
        return config.TABLA_NOTA_ASISTENCIA_9_CLASES.get(clases_asistidas, 0)
    # Curso con un número de clases distinto de 9: escala proporcional.
    nota = round((clases_asistidas / curso.num_clases) * config.NOTA_MAXIMA_ASISTENCIA)
    return max(0, min(config.NOTA_MAXIMA_ASISTENCIA, nota))


def _aprobo_por_asistencia(clases_asistidas: int) -> bool:
    return clases_asistidas >= config.NUM_CLASES_MINIMAS_PARA_APROBAR


def _fue_entregado(valor_crudo) -> bool:
    """True si la celda representa una entrega/nota real (no vacía, no
    un placeholder tipo '-', y distinta de cero)."""
    if valor_crudo in (None, "", "-"):
        return False
    return _a_entero(valor_crudo) != 0


# ============================================================
# Paso 3: procesar los archivos de cada tutora
# ============================================================

def _procesar_archivo_tutor(
    archivo: Path, ws_main: Worksheet, dni_index: dict[str, int],
    idx: dict, curso: ConfiguracionCurso, resultado: ResultadoCalculo,
) -> None:
    try:
        wb_tutor = openpyxl.load_workbook(archivo, data_only=True, read_only=True)
    except Exception as e:
        raise ReporteFinalError(f"No se pudo abrir el archivo de la tutora '{archivo.name}': {e}") from e

    if config.NOMBRE_HOJA_ASISTENCIA_TUTOR not in wb_tutor.sheetnames:
        resultado.archivos_tutor_sin_hoja_asistencia.append(archivo.name)
        logger.warning("Archivo %s sin hoja %s", archivo, config.NOMBRE_HOJA_ASISTENCIA_TUTOR)
        return

    ws_asistencia = wb_tutor[config.NOMBRE_HOJA_ASISTENCIA_TUTOR]
    ws_devo = wb_tutor[config.NOMBRE_HOJA_DEVOCIONALES_TUTOR] if config.NOMBRE_HOJA_DEVOCIONALES_TUTOR in wb_tutor.sheetnames else None

    # Nombre de archivo tal cual, solo reemplazando guiones bajos por
    # espacios — así la columna "Grupo" queda 100% trazable al archivo
    # real, sin adivinar qué parte del nombre es "la tutora".
    tutor_nombre = archivo.stem.replace("_", " ")
    resultado.tutores_procesados.append(tutor_nombre)
    logger.info("Procesando tutora: %s (%s)", tutor_nombre, archivo.name)

    col_dni_t = config.TUTOR_COL_DNI - 1              # a índice 0
    col_asist_ini_t = config.TUTOR_COL_ASISTENCIA_INICIO - 1
    col_tf_t = config.TUTOR_COL_TF_DRIVE - 1
    col_ef_t = config.TUTOR_COL_EF_DRIVE - 1
    col_devo_ini_t = config.TUTOR_DEVO_COL_INICIO - 1

    for row_idx, row in enumerate(
        ws_asistencia.iter_rows(min_row=config.TUTOR_FILA_INICIO_DATOS),
        start=config.TUTOR_FILA_INICIO_DATOS,
    ):
        dni_celda = row[col_dni_t].value if len(row) > col_dni_t else None
        if not dni_celda:
            break  # fin de la lista de alumnos de esta tutora

        dni_tutor = normalizar_dni_valor(dni_celda)
        if dni_tutor not in dni_index:
            resultado.dnis_no_encontrados_en_reporte.append(dni_tutor)
            continue

        r = dni_index[dni_tutor]

        # El DNI/carnet que trae Moodle puede tener ceros a la izquierda
        # perdidos (Moodle lo exporta como número). El de la tutora sí
        # está escrito como texto, con el formato correcto — se usa ese
        # para lo que se muestra en el reporte final.
        dni_texto_tutor = str(dni_celda).strip()
        if dni_texto_tutor:
            ws_main.cell(r, idx["dni"], dni_texto_tutor)

        # "Grupo" muestra de qué archivo de tutora salió este alumno,
        # para poder rastrear a qué tutora pertenece cada fila.
        ws_main.cell(r, idx["grupo"], tutor_nombre)

        asistencias = row[col_asist_ini_t: col_asist_ini_t + curso.num_clases]
        total_asist = contar_celdas_con_valor(asistencias)
        nota_asist = calcular_nota_asistencia(total_asist, curso)

        # -------- copiar asistencia por clase (columnas S1..Sn) --------
        for j, celda in enumerate(asistencias):
            target = ws_main.cell(r, idx["clases_drive_inicio"] + j, celda.value)
            set_fill_por_valor(celda.value, target)

        ws_main.cell(r, idx["clases_drive"], total_asist)
        ws_main.cell(r, idx["asistencia_drive"], nota_asist)

        if total_asist == curso.num_clases:
            ws_main.cell(r, idx["clases_drive"]).fill = PatternFill("solid", fgColor=config.COLOR_ASISTENCIA_COMPLETA)
        elif total_asist / curso.num_clases >= config.UMBRAL_ASISTENCIA_PARCIAL:
            ws_main.cell(r, idx["clases_drive"]).fill = PatternFill("solid", fgColor=config.COLOR_ASISTENCIA_PARCIAL)

        # -------- comparar asistencia Drive vs AV --------
        val_av = _a_entero(ws_main.cell(r, idx["asistencia_av"]).value)
        if val_av != nota_asist:
            ws_main.cell(r, idx["asistencia_av"]).fill = PatternFill("solid", fgColor=config.COLOR_CELDA_INCONSISTENCIA)

        # -------- Examen Final --------
        # "indicadores" junta las notas (0-20) que sí aplican a este curso,
        # para el promedio final. Los "gates" (ef_resuelto / tf_entregado)
        # son condiciones aparte: aunque el promedio dé bien, si el curso
        # tiene examen y no lo resolvió (o tiene TF y no lo entregó), no
        # aprueba igual.
        indicadores = [nota_asist]
        ef_resuelto = True
        if curso.tiene_examen_final and idx["ef_drive"]:
            ef_drive_raw = row[col_ef_t].value if len(row) > col_ef_t else None
            ef_drive_val = _a_entero(ef_drive_raw)
            ws_main.cell(r, idx["ef_drive"], ef_drive_val)
            if idx["ef_av"]:
                val_ef_av = _a_entero(ws_main.cell(r, idx["ef_av"]).value)
                if val_ef_av != ef_drive_val:
                    ws_main.cell(r, idx["ef_drive"]).fill = PatternFill("solid", fgColor=config.COLOR_CELDA_INCONSISTENCIA)
            ef_resuelto = _fue_entregado(ef_drive_raw)
            indicadores.append(ef_drive_val)

        # -------- Trabajo Final --------
        tf_entregado = True
        if curso.tiene_trabajo_final and idx["tf_drive"]:
            tf_drive_raw = row[col_tf_t].value if len(row) > col_tf_t else None
            tf_drive_val = _a_entero(tf_drive_raw)
            ws_main.cell(r, idx["tf_drive"], tf_drive_val)
            if idx["tf_av"]:
                val_tf_av = _a_entero(ws_main.cell(r, idx["tf_av"]).value)
                if val_tf_av != tf_drive_val:
                    ws_main.cell(r, idx["tf_drive"]).fill = PatternFill("solid", fgColor=config.COLOR_CELDA_INCONSISTENCIA)
            tf_entregado = _fue_entregado(tf_drive_raw)
            indicadores.append(tf_drive_val)

        # -------- Devocionales (siempre, para todo curso) --------
        if ws_devo is not None:
            devo_row_idx = row_idx + config.TUTOR_DEVO_FILA_OFFSET
            devo_row = ws_devo[devo_row_idx]
            devo_celdas = devo_row[col_devo_ini_t: col_devo_ini_t + config.TUTOR_DEVO_TOTAL_DIAS]
            total_devo = sum(1 for c in devo_celdas if c.value == 1)
            pct_devo = round((total_devo / config.TUTOR_DEVO_TOTAL_DIAS) * 100, 2) if config.TUTOR_DEVO_TOTAL_DIAS else 0
            ws_main.cell(r, idx["devo_total"], total_devo)
            ws_main.cell(r, idx["devo_pct"], pct_devo)

        # -------- estado final: APROBADO / DESAPROBADO --------
        # 1) mínimo de asistencia, 2) resolvió el examen (si aplica),
        # 3) entregó el trabajo final (si aplica), 4) promedio >= 10.5.
        promedio = sum(indicadores) / len(indicadores)
        aprobo = (
            _aprobo_por_asistencia(total_asist)
            and ef_resuelto
            and tf_entregado
            and promedio >= config.PROMEDIO_MINIMO_PARA_APROBAR
        )
        celda_estado = ws_main.cell(r, idx["estado_final"], "APROBADO" if aprobo else "DESAPROBADO")
        if aprobo:
            celda_estado.fill = PatternFill("solid", fgColor=config.COLOR_APROBADO_FONDO)
            celda_estado.font = Font(color=config.COLOR_APROBADO_TEXTO, bold=True)
        else:
            celda_estado.fill = PatternFill("solid", fgColor=config.COLOR_DESAPROBADO_FONDO)
            celda_estado.font = Font(color=config.COLOR_DESAPROBADO_TEXTO, bold=True)
        celda_estado.alignment = Alignment(horizontal="center")


# ============================================================
# Paso 4: comparar checks (Clase N) de Moodle vs asistencia del Drive
# ============================================================

def _marcar_inconsistencias_clase_por_clase(ws_main: Worksheet, idx: dict, curso: ConfiguracionCurso) -> None:
    mapa = idx["mapa"]
    columnas_estilizadas = idx["columnas_estilizadas"]
    col_falta_check = _agregar_columna(ws_main, "Falta check", mapa, columnas_estilizadas)
    col_falta_apreciacion = _agregar_columna(ws_main, "Falta apreciación", mapa, columnas_estilizadas)
    ws_main.cell(1, col_falta_check).fill = PatternFill("solid", fgColor=config.COLOR_NEGRO)
    ws_main.cell(1, col_falta_check).font = Font(bold=True, color=config.COLOR_AMARILLO)
    ws_main.cell(1, col_falta_apreciacion).fill = PatternFill("solid", fgColor=config.COLOR_ROJO)
    ws_main.cell(1, col_falta_apreciacion).font = Font(bold=True, color=config.COLOR_BLANCO)

    columnas_clase = []
    for j in range(1, curso.num_clases + 1):
        col_clase = mapa.get(f"Clase {j}")
        col_s = mapa.get(f"S{j}")
        if col_clase and col_s:
            columnas_clase.append((j, col_clase, col_s))
        else:
            logger.warning("No se encontró columna 'Clase %d' o 'S%d' para el cruce de checks.", j, j)

    for r in range(2, ws_main.max_row + 1):
        faltan_check, faltan_apreciacion = [], []
        for j, col_clase, col_s in columnas_clase:
            estado_av = ws_main.cell(r, col_clase).value
            valor_drive = ws_main.cell(r, col_s).value
            celda_clase = ws_main.cell(r, col_clase)

            if estado_av == "No finalizado":
                celda_clase.fill = PatternFill("solid", fgColor=config.COLOR_NEGRO)
                celda_clase.font = Font(color=config.COLOR_AMARILLO if valor_drive else config.COLOR_BLANCO)
                if valor_drive:
                    faltan_check.append(str(j))
            elif estado_av == "Finalizado" and not valor_drive:
                celda_clase.fill = PatternFill("solid", fgColor=config.COLOR_ROJO)
                celda_clase.font = Font(color=config.COLOR_BLANCO)
                faltan_apreciacion.append(str(j))

        def _texto(lista):
            if not lista:
                return ""
            return ("Clases " if len(lista) > 1 else "Clase ") + ", ".join(lista)

        ws_main.cell(r, col_falta_check, _texto(faltan_check))
        ws_main.cell(r, col_falta_apreciacion, _texto(faltan_apreciacion))


def _aplicar_formato_general(ws_main: Worksheet, columnas_estilizadas: set[int]) -> None:
    """Aplica el estilo azul/blanco por defecto SOLO a los encabezados que
    vinieron tal cual de Moodle (Nombre, DNI, Asistencia AV, etc.) y que
    nunca pasaron por _agregar_columna(). Las columnas que sí pasaron por
    ahí (incluidas las que después se recolorearon a mano, como "Falta
    check" en negro/amarillo) están en `columnas_estilizadas` y se dejan
    tal cual — no se detectan por su color actual, porque un color como
    negro puro es indistinguible de "sin color" para openpyxl y terminaba
    pisándose con el azul por defecto."""
    for col in range(1, ws_main.max_column + 1):
        cell = ws_main.cell(1, col)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if col in columnas_estilizadas:
            continue
        cell.font = Font(bold=True, color=config.COLOR_BLANCO)
        cell.fill = PatternFill("solid", fgColor=config.COLOR_AZUL_HEADER)

    for fila in ws_main.iter_rows(min_row=1, max_row=ws_main.max_row, min_col=1, max_col=ws_main.max_column):
        for cell in fila:
            cell.border = BORDE_FINO


# ============================================================
# Punto de entrada
# ============================================================

def generar_reporte_final(
    reporte_unido_path: Path, tutores_dir: Path, out_path: Path, curso: ConfiguracionCurso,
) -> tuple[Path, ResultadoCalculo]:
    try:
        wb_main = openpyxl.load_workbook(reporte_unido_path)
        ws_main = wb_main.active
    except Exception as e:
        raise ReporteFinalError(f"No se pudo abrir el reporte unido: {e}") from e

    resultado = ResultadoCalculo()
    idx = _preparar_columnas(ws_main, curso)

    dni_index: dict[str, int] = {}
    for r in range(2, ws_main.max_row + 1):
        val = ws_main.cell(r, idx["dni"]).value
        if val:
            dni_index[normalizar_dni_valor(val)] = r
    logger.info("Índice de DNI cargado: %d registros", len(dni_index))

    archivos_tutores = sorted(Path(tutores_dir).glob("*.xlsx"))
    if not archivos_tutores:
        raise ReporteFinalError(f"No se encontró ningún archivo .xlsx de tutora en {tutores_dir}.")

    for archivo in archivos_tutores:
        _procesar_archivo_tutor(archivo, ws_main, dni_index, idx, curso, resultado)

    if resultado.dnis_no_encontrados_en_reporte:
        logger.warning(
            "%d DNI(s) de archivos de tutora no se encontraron en el reporte: %s",
            len(resultado.dnis_no_encontrados_en_reporte),
            sorted(set(resultado.dnis_no_encontrados_en_reporte)),
        )

    _marcar_inconsistencias_clase_por_clase(ws_main, idx, curso)
    _aplicar_formato_general(ws_main, idx["columnas_estilizadas"])
    autoajustar_columnas(ws_main)

    wb_main.save(out_path)
    logger.info("Archivo final generado: %s", out_path)
    return out_path, resultado