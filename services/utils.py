"""Utilidades genéricas de formato y manejo de DataFrames/Excel."""
from __future__ import annotations

from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
import pandas as pd

import config


def normalizar_dni_valor(valor) -> str:
    """Normaliza un identificador de alumno (DNI o carnet de extranjería)
    para poder CRUZARLO entre Moodle y el archivo de la tutora.

    Moodle exporta este campo como NÚMERO tanto en el Excel de notas
    como en el CSV de checks, así que cualquier cero a la izquierda se
    pierde ahí (ej: DNI "07144588" -> 7144588, carnet "003112762" -> 3112762).
    El archivo de la tutora, en cambio, sí guarda el valor como texto,
    con los ceros.

    Como el DNI siempre tiene 8 dígitos, pero el carnet de extranjería NO
    tiene un largo fijo (puede ser de 9 o más), no se puede "adivinar"
    cuántos ceros había forzando un ancho fijo — eso funciona para DNI,
    pero rompe carnets de extranjería. En su lugar, para cruzar ambos
    lados se quitan TODOS los ceros a la izquierda de los dos por igual:
    así "07144588" y "7144588" cruzan entre sí, y también "003112762"
    con "3112762", sin importar el largo real de cada uno.

    Esto es solo para USAR COMO LLAVE DE CRUCE, no para mostrar en el
    reporte — para eso, se prefiere el valor tal como está escrito en el
    archivo de la tutora (ver services/calculator.py), que si está en
    texto ya tiene el formato/ceros correctos.
    """
    texto = str(valor).strip()
    if texto.endswith(".0"):
        texto = texto[:-2]
    return texto.lstrip("0") or "0"


def normalizar_dni(serie: pd.Series) -> pd.Series:
    """Igual que normalizar_dni_valor(), pero aplicado a una columna completa."""
    return serie.apply(normalizar_dni_valor)


def autoajustar_columnas(ws: Worksheet) -> None:
    """Ajusta el ancho de cada columna al contenido más largo que tenga."""
    for col in ws.columns:
        max_length = 0
        column_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[column_letter].width = max_length + 2


def set_fill_por_valor(value, target_cell) -> None:
    """Colorea una celda de asistencia: verde si hay un valor real
    (el alumno asistió), gris si está vacía (no asistió)."""
    if value not in (0, None, ""):
        color = config.COLOR_CHECK_PRESENTE
    else:
        color = config.COLOR_CHECK_AUSENTE
    target_cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")


def contar_celdas_con_valor(celdas) -> int:
    """Cuenta cuántas celdas de una fila de asistencia tienen un valor > 0."""
    total = 0
    for c in celdas:
        val = c.value
        if val is None:
            continue
        try:
            if float(val) > 0:
                total += 1
        except (ValueError, TypeError):
            pass
    return total