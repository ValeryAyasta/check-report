from pathlib import Path

import openpyxl
import pytest

from domain.curso import ConfiguracionCurso
from services.calculator import generar_reporte_final, calcular_nota_asistencia


def _reporte_unido_xlsx(path: Path, alumnos: list[dict], con_tf: bool, con_ef: bool) -> None:
    """alumnos: [{"nombre","dni","tf_av","ef_av"}]"""
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = ["Nombre", "Apellidos", "DNI", "Asistencia AV"]
    if con_tf:
        headers.append("TF (AV)")
    if con_ef:
        headers.append("EF (AV)")
    ws.append(headers)
    for a in alumnos:
        fila = [a["nombre"], "Apellido", a["dni"], 0]
        if con_tf:
            fila.append(a.get("tf_av", 0))
        if con_ef:
            fila.append(a.get("ef_av", 0))
        ws.append(fila)
    wb.save(path)


def _tutor_xlsx(path: Path, alumnos: list[dict], num_clases: int = 9) -> None:
    """alumnos: [{"dni","clases","tf_drive","ef_drive"}]"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ASISTENCIA"
    for _ in range(3):
        ws.append([None] * 40)
    for a in alumnos:
        fila = [None] * 4 + [a["dni"]]
        fila += [1] * a["clases"] + [0] * (num_clases - a["clases"])
        fila += [None] * (34 - len(fila))   # relleno hasta el índice 33 (columna AH)
        fila.append(a.get("tf_drive", 0))   # índice 34 = columna 35 (AI) = TF
        fila.append(None)                   # índice 35 = columna 36 (AJ)
        fila.append(a.get("ef_drive", 0))   # índice 36 = columna 37 (AK) = EF
        ws.append(fila)
    ws_devo = wb.create_sheet("DEVO")
    for _ in range(3 + len(alumnos) + 1):
        ws_devo.append([None] * 70)
    wb.save(path)


CURSO_NORMAL = ConfiguracionCurso(curso_id=1, nombre="Curso", grupos_texto="1", num_clases=9)


class TestFormulaDeAprobacion:
    """Los 5 escenarios validados manualmente al implementar la regla:
    mínimo 6 clases AND resolvió examen (si aplica) AND entregó TF (si
    aplica) AND promedio(indicadores aplicables) >= 10.5."""

    @pytest.mark.parametrize("caso,esperado", [
        pytest.param(
            {"clases": 9, "ef_drive": 18, "tf_drive": 18}, "APROBADO",
            id="notas_altas_9_clases",
        ),
        pytest.param(
            {"clases": 5, "ef_drive": 20, "tf_drive": 20}, "DESAPROBADO",
            id="asiste_5_no_llega_al_minimo_de_6",
        ),
        pytest.param(
            {"clases": 6, "ef_drive": 5, "tf_drive": 5}, "DESAPROBADO",
            id="cumple_minimos_pero_promedio_bajo",
        ),
        pytest.param(
            {"clases": 9, "ef_drive": 0, "tf_drive": 18}, "DESAPROBADO",
            id="no_resuelve_examen",
        ),
        pytest.param(
            {"clases": 6, "ef_drive": 8, "tf_drive": 13}, "APROBADO",
            id="promedio_justo_arriba_de_10_5",  # asist=13 (tabla 6->13); (13+8+13)/3=11.33
        ),
    ])
    def test_escenario(self, tmp_path, caso, esperado):
        alumno = {"nombre": "Alumno", "dni": "10000001", "tf_av": caso["tf_drive"], "ef_av": caso["ef_drive"]}
        tutor = {"dni": "10000001", **caso}

        ruta_reporte = tmp_path / "reporte.xlsx"
        ruta_tutores = tmp_path / "tutores"
        ruta_tutores.mkdir()
        _reporte_unido_xlsx(ruta_reporte, [alumno], con_tf=True, con_ef=True)
        _tutor_xlsx(ruta_tutores / "Tutora.xlsx", [tutor])

        out_path, _ = generar_reporte_final(ruta_reporte, ruta_tutores, tmp_path / "final.xlsx", CURSO_NORMAL)

        wb = openpyxl.load_workbook(out_path)
        ws = wb.active
        mapa = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
        assert ws.cell(2, mapa["Estado Final"]).value == esperado

    def test_curso_sin_examen_ni_trabajo_final_solo_usa_asistencia(self, tmp_path):
        curso = ConfiguracionCurso(curso_id=1, nombre="Curso", grupos_texto="1", num_clases=9,
                                    sin_examen_final=True, sin_trabajo_final=True)
        alumno = {"nombre": "Alumno", "dni": "10000002"}
        tutor = {"dni": "10000002", "clases": 9}

        ruta_reporte = tmp_path / "reporte.xlsx"
        ruta_tutores = tmp_path / "tutores"
        ruta_tutores.mkdir()
        _reporte_unido_xlsx(ruta_reporte, [alumno], con_tf=False, con_ef=False)
        _tutor_xlsx(ruta_tutores / "Tutora.xlsx", [tutor])

        out_path, _ = generar_reporte_final(ruta_reporte, ruta_tutores, tmp_path / "final.xlsx", curso)
        wb = openpyxl.load_workbook(out_path)
        ws = wb.active
        mapa = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
        assert ws.cell(2, mapa["Estado Final"]).value == "APROBADO"  # asistencia=20, único indicador
        assert "EF Drive" not in mapa
        assert "TF Drive" not in mapa
        # Devocionales SIEMPRE se calculan, incluso sin examen/TF:
        assert "Devocionales Entregados" in mapa

    def test_devocionales_siempre_presentes(self, tmp_path):
        """Confirma que ya no existe la excepción 'curso sin trabajo del
        libro': las columnas de devocionales están en todo reporte."""
        alumno = {"nombre": "Alumno", "dni": "10000003", "tf_av": 15, "ef_av": 15}
        tutor = {"dni": "10000003", "clases": 9, "tf_drive": 15, "ef_drive": 15}

        ruta_reporte = tmp_path / "reporte.xlsx"
        ruta_tutores = tmp_path / "tutores"
        ruta_tutores.mkdir()
        _reporte_unido_xlsx(ruta_reporte, [alumno], con_tf=True, con_ef=True)
        _tutor_xlsx(ruta_tutores / "Tutora.xlsx", [tutor])

        out_path, _ = generar_reporte_final(ruta_reporte, ruta_tutores, tmp_path / "final.xlsx", CURSO_NORMAL)
        wb = openpyxl.load_workbook(out_path)
        ws = wb.active
        mapa = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
        assert "Devocionales Entregados" in mapa
        assert "% Devocionales" in mapa


class TestCruceDeIdentificadores:
    def test_dni_con_cero_a_la_izquierda_cruza_bien(self, tmp_path):
        alumno = {"nombre": "Ana", "dni": "7144588"}  # como llega de Moodle, ya sin el cero
        tutor = {"dni": "07144588", "clases": 9}  # como está en el Drive, con el cero

        ruta_reporte = tmp_path / "reporte.xlsx"
        ruta_tutores = tmp_path / "tutores"
        ruta_tutores.mkdir()
        _reporte_unido_xlsx(ruta_reporte, [alumno], con_tf=False, con_ef=False)
        _tutor_xlsx(ruta_tutores / "Tutora.xlsx", [tutor])

        out_path, resultado = generar_reporte_final(
            ruta_reporte, ruta_tutores, tmp_path / "final.xlsx",
            ConfiguracionCurso(curso_id=1, nombre="X", grupos_texto="1", num_clases=9,
                                sin_examen_final=True, sin_trabajo_final=True),
        )
        assert resultado.dnis_no_encontrados_en_drive_pero_no_en_moodle == []

        wb = openpyxl.load_workbook(out_path)
        ws = wb.active
        mapa = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
        # El reporte final debe mostrar el DNI completo (tomado del Drive), no el recortado.
        assert ws.cell(2, mapa["DNI"]).value == "07144588"

    def test_carnet_extranjeria_9_digitos_cruza_bien(self, tmp_path):
        """El carnet de extranjería no tiene largo fijo de 8 dígitos como
        el DNI — no se le debe forzar zfill(8)."""
        alumno = {"nombre": "Ana", "dni": "3112762"}  # Moodle perdió 2 ceros
        tutor = {"dni": "003112762", "clases": 9}  # Drive: 9 dígitos reales

        ruta_reporte = tmp_path / "reporte.xlsx"
        ruta_tutores = tmp_path / "tutores"
        ruta_tutores.mkdir()
        _reporte_unido_xlsx(ruta_reporte, [alumno], con_tf=False, con_ef=False)
        _tutor_xlsx(ruta_tutores / "Tutora.xlsx", [tutor])

        out_path, resultado = generar_reporte_final(
            ruta_reporte, ruta_tutores, tmp_path / "final.xlsx",
            ConfiguracionCurso(curso_id=1, nombre="X", grupos_texto="1", num_clases=9,
                                sin_examen_final=True, sin_trabajo_final=True),
        )
        assert resultado.dnis_no_encontrados_en_drive_pero_no_en_moodle == []

        wb = openpyxl.load_workbook(out_path)
        ws = wb.active
        mapa = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
        assert ws.cell(2, mapa["DNI"]).value == "003112762"

    def test_grupo_se_llena_con_nombre_de_archivo_de_tutora(self, tmp_path):
        alumno = {"nombre": "Ana", "dni": "10000004"}
        tutor = {"dni": "10000004", "clases": 9}

        ruta_reporte = tmp_path / "reporte.xlsx"
        ruta_tutores = tmp_path / "tutores"
        ruta_tutores.mkdir()
        _reporte_unido_xlsx(ruta_reporte, [alumno], con_tf=False, con_ef=False)
        _tutor_xlsx(ruta_tutores / "Equipo1-Maria_Lopez.xlsx", [tutor])

        curso = ConfiguracionCurso(curso_id=1, nombre="X", grupos_texto="1", num_clases=9,
                                    sin_examen_final=True, sin_trabajo_final=True)
        out_path, _ = generar_reporte_final(ruta_reporte, ruta_tutores, tmp_path / "final.xlsx", curso)

        wb = openpyxl.load_workbook(out_path)
        ws = wb.active
        mapa = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
        assert ws.cell(2, mapa["Grupo"]).value == "Equipo1-Maria Lopez"


class TestNotaDeAsistencia:
    @pytest.mark.parametrize("asistidas,esperado", [
        (0, 0), (1, 2), (5, 10), (6, 13), (9, 20),
    ])
    def test_tabla_de_9_clases(self, asistidas, esperado):
        curso = ConfiguracionCurso(curso_id=1, nombre="X", grupos_texto="1", num_clases=9)
        assert calcular_nota_asistencia(asistidas, curso) == esperado

    def test_curso_de_8_clases_usa_escala_proporcional(self):
        curso = ConfiguracionCurso(curso_id=1, nombre="X", grupos_texto="1", num_clases=8)
        assert calcular_nota_asistencia(8, curso) == 20  # asistió a todas
        assert calcular_nota_asistencia(0, curso) == 0
        assert calcular_nota_asistencia(4, curso) == 10  # la mitad


class TestAdvertenciasDeCruce:
    def test_alumno_en_moodle_sin_registro_en_drive_se_reporta(self, tmp_path):
        """El alumno existe en el reporte de Moodle pero ninguna tutora
        tiene su DNI en su archivo de Drive."""
        alumnos = [
            {"nombre": "Ana", "dni": "10000010"},
            {"nombre": "Beto", "dni": "10000011"},  # este NO va a estar en ningún archivo de tutora
        ]
        ruta_reporte = tmp_path / "reporte.xlsx"
        ruta_tutores = tmp_path / "tutores"
        ruta_tutores.mkdir()
        _reporte_unido_xlsx(ruta_reporte, alumnos, con_tf=False, con_ef=False)
        _tutor_xlsx(ruta_tutores / "Tutora.xlsx", [{"dni": "10000010", "clases": 9}])  # solo Ana

        curso = ConfiguracionCurso(curso_id=1, nombre="X", grupos_texto="C11", num_clases=9,
                                    sin_examen_final=True, sin_trabajo_final=True)
        out_path, resultado = generar_reporte_final(ruta_reporte, ruta_tutores, tmp_path / "final.xlsx", curso)

        assert len(resultado.alumnos_en_moodle_sin_registro_en_drive) == 1
        assert "10000011" in resultado.alumnos_en_moodle_sin_registro_en_drive[0]
        assert resultado.dnis_no_encontrados_en_drive_pero_no_en_moodle == []

    def test_dni_en_drive_pero_no_en_moodle_se_reporta_por_separado(self, tmp_path):
        """Caso opuesto: el DNI está en un archivo de tutora, pero no
        existe en el reporte de Moodle (otro grupo/curso, o mal tipeado)."""
        alumnos = [{"nombre": "Ana", "dni": "10000020"}]
        ruta_reporte = tmp_path / "reporte.xlsx"
        ruta_tutores = tmp_path / "tutores"
        ruta_tutores.mkdir()
        _reporte_unido_xlsx(ruta_reporte, alumnos, con_tf=False, con_ef=False)
        _tutor_xlsx(ruta_tutores / "Tutora.xlsx", [
            {"dni": "10000020", "clases": 9},
            {"dni": "99999999", "clases": 9},  # no está en el reporte de Moodle
        ])

        curso = ConfiguracionCurso(curso_id=1, nombre="X", grupos_texto="C11", num_clases=9,
                                    sin_examen_final=True, sin_trabajo_final=True)
        out_path, resultado = generar_reporte_final(ruta_reporte, ruta_tutores, tmp_path / "final.xlsx", curso)

        assert resultado.dnis_no_encontrados_en_drive_pero_no_en_moodle == ["99999999"]
        assert resultado.alumnos_en_moodle_sin_registro_en_drive == []


class TestExtraerDnisTutor:
    def test_extrae_los_dnis_de_un_archivo(self, tmp_path):
        from services.calculator import extraer_dnis_tutor
        ruta = tmp_path / "Tutora.xlsx"
        _tutor_xlsx(ruta, [{"dni": "10000030", "clases": 9}, {"dni": "07144588", "clases": 5}])

        dnis = extraer_dnis_tutor(ruta)
        assert dnis == {"10000030", "7144588"}  # normalizado, sin ceros a la izquierda

    def test_archivo_sin_hoja_asistencia_devuelve_vacio(self, tmp_path):
        import openpyxl
        from services.calculator import extraer_dnis_tutor
        ruta = tmp_path / "SinHoja.xlsx"
        openpyxl.Workbook().save(ruta)  # hoja por defecto "Sheet", no "ASISTENCIA"

        assert extraer_dnis_tutor(ruta) == set()
