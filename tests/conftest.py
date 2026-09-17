"""
Fixtures compartidas. Generan archivos Excel/CSV EN MEMORIA con la misma
estructura que se verificó contra archivos reales de Moodle y de una
tutora (ver el hilo de la conversación donde se construyó este
proyecto) — pero con datos 100% inventados, para no depender ni
commitear ningún dato real de alumnos.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))


def _construir_notas_xlsx(alumnos: list[dict], con_tf: bool, con_ef: bool) -> bytes:
    """
    alumnos: [{"nombre": ..., "apellidos": ..., "dni": ..., "grupo": ...,
               "asistencia": int, "tf": int, "ef": int}]
    Replica el layout REAL verificado del grade export de Moodle.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = ["Nombre", "Apellido(s)", "Nombre de usuario", "Dirección de correo",
               "Teléfono", "Grupo", "Asistencia (Real)"]
    if con_tf:
        headers.append("Tarea:Tarea Final (Real)")
    if con_ef:
        headers.append("Cuestionario:Examen Final (Real)")
    headers += ["Total del curso (Real)", "Última descarga de este curso"]
    ws.append(headers)

    for a in alumnos:
        fila = [a["nombre"], a["apellidos"], int(a["dni"]), f"{a['dni']}@correo.com",
                "999999999", a.get("grupo", "C11"), a["asistencia"]]
        if con_tf:
            fila.append(a.get("tf", 0))
        if con_ef:
            fila.append(a.get("ef", 0))
        fila += [0, "-"]
        ws.append(fila)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _construir_checks_csv(alumnos: list[dict], num_clases: int) -> bytes:
    """Replica el layout REAL del CSV de checks (tab-separated, UTF-16-LE,
    con columnas 'Clase N' + una columna de fecha por cada una)."""
    columnas = ["Nombre completo", "Nombre de usuario", "Dirección de correo"]
    for i in range(1, num_clases + 1):
        columnas += [f"Clase {i}", ""]  # la columna de fecha va sin encabezado en el archivo real

    lineas = ["\t".join(columnas)]
    for a in alumnos:
        fila = [f"{a['nombre']} {a['apellidos']}", str(int(a["dni"])), f"{a['dni']}@correo.com"]
        clases_completadas = a.get("clases_finalizadas", num_clases)
        for i in range(1, num_clases + 1):
            estado = "Finalizado" if i <= clases_completadas else "No finalizado"
            fecha = "2026-01-01" if i <= clases_completadas else ""
            fila += [estado, fecha]
        lineas.append("\t".join(fila))

    texto = "\n".join(lineas)
    return texto.encode("utf-16-le")


def _construir_tutor_xlsx(alumnos: list[dict], num_clases: int = 9) -> bytes:
    """
    alumnos: [{"dni": str (con formato real, ej '07144588'), "clases": int,
               "tf_drive": int, "ef_drive": int, "devo_dias": int}]
    Replica el layout REAL verificado del archivo de asistencia de una
    tutora (hojas ASISTENCIA y DEVO).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ASISTENCIA"
    for _ in range(3):
        ws.append([None] * 40)

    for a in alumnos:
        fila = [None] * 4 + [a["dni"]]                       # A-D vacío, E = DNI
        fila += [1] * a["clases"] + [0] * (num_clases - a["clases"])  # F..N = clases
        fila += [None] * (34 - len(fila))                    # relleno hasta índice 33 (col AH)
        fila.append(a.get("tf_drive", 0))                    # índice 34 = col 35 (AI) = TF
        fila.append(None)                                     # índice 35 = col 36 (AJ)
        fila.append(a.get("ef_drive", 0))                    # índice 36 = col 37 (AK) = EF
        ws.append(fila)

    ws_devo = wb.create_sheet("DEVO")
    for _ in range(4):
        ws_devo.append([None] * 70)
    for a in alumnos:
        dias = a.get("devo_dias", 0)
        fila = [None] * 6 + [1] * dias + [0] * (63 - dias)   # G..BQ = 63 días
        ws_devo.append(fila)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def construir_notas_xlsx():
    return _construir_notas_xlsx


@pytest.fixture
def construir_checks_csv():
    return _construir_checks_csv


@pytest.fixture
def construir_tutor_xlsx():
    return _construir_tutor_xlsx
